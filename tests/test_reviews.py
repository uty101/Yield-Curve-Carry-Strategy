"""Step 0.3: every review file follows review/TEMPLATE.md; the fixture trimmer keeps 50 rows."""

import importlib.util
import re
from pathlib import Path

import pytest

REVIEW_DIR = Path("review")
HEADINGS = ["## What changed", "## Findings", "## Not verified", "## Open", "## Reviewer reads"]
FIRST_LINE = re.compile(r"^# Review — Step \d+\.\d+ — .+")
META_LINE = re.compile(r"^Date: .+Commit: .+Tests: \d+/\d+")


def _review_files() -> list[Path]:
    return sorted(p for p in REVIEW_DIR.glob("*.md") if p.name != "TEMPLATE.md")


def _claim_blocks(lines: list[str]) -> list[list[str]]:
    """Each block: the lines from a '> Claim:' line to the next blank line."""
    blocks, cur = [], None
    for line in lines:
        if line.startswith("> Claim:"):
            cur = [line]
        elif cur is not None:
            if line.strip() == "":
                blocks.append(cur)
                cur = None
            else:
                cur.append(line)
    if cur is not None:
        blocks.append(cur)
    return blocks


@pytest.mark.parametrize("path", _review_files() or [None])
def test_reviews_follow_template(path: Path | None) -> None:
    if path is None:  # no reviews yet: passes vacuously
        return
    lines = path.read_text(encoding="utf-8").splitlines()
    assert FIRST_LINE.match(lines[0]), f"{path}: first line"
    assert any(META_LINE.match(ln) for ln in lines[:5]), f"{path}: Date/Commit/Tests line"
    for h in HEADINGS:
        assert any(ln.startswith(h) for ln in lines), f"{path}: missing {h!r}"
    blocks = _claim_blocks(lines)
    assert blocks, f"{path}: no '> Claim:' block"
    for b in blocks:
        assert len(b) > 1 and b[1].startswith("> Number:"), f"{path}: {b[0][:60]!r} lacks Number"
        rows = [ln for ln in b[2:] if ln.startswith("> ") and not ln.startswith("> Rows")]
        assert len(rows) >= 5, f"{path}: {b[0][:60]!r} has {len(rows)} rows, needs 5"


def _load_make_fixture():
    spec = importlib.util.spec_from_file_location("make_fixture", "scripts/make_fixture.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_make_fixture_trims_to_50_rows(tmp_path: Path) -> None:
    import openpyxl

    mf = _load_make_fixture()
    raw_csv = tmp_path / "raw" / "DGS10_20260916.csv"
    raw_csv.parent.mkdir()
    raw_csv.write_text(
        "observation_date,DGS10\n" + "".join(f"2000-01-{i:02d},4.{i:02d}\n" for i in range(1, 201)),
        encoding="utf-8",
    )
    out_csv = mf.make_fixture("fred", raw_csv, fixtures=tmp_path / "fixtures")
    assert out_csv == tmp_path / "fixtures" / "fred" / "DGS10.csv"
    lines = out_csv.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "observation_date,DGS10"
    assert len(lines) == 51 and lines[1] == "2000-01-01,4.01" and lines[50] == "2000-01-50,4.50"

    raw_x = tmp_path / "raw" / "GLC Nominal month end data_2016 to 2024.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "1. fwds"
    ws = wb.create_sheet("4. spot curve")
    ws.append(["years:", 0.5, 1.0, 1.5])
    for i in range(200):
        ws.append([f"row{i}", 4.0 + i, 4.1 + i, 4.2 + i])
    wb.save(raw_x)
    out_x = mf.make_fixture("boe", raw_x, sheet="4. spot curve", fixtures=tmp_path / "fixtures")
    got = openpyxl.load_workbook(out_x)
    assert got.sheetnames == ["4. spot curve"]
    rows = list(got["4. spot curve"].iter_rows(values_only=True))
    assert len(rows) == 51
    assert rows[0][0] == "years:" and rows[1][0] == "row0" and rows[50][0] == "row49"

    with pytest.raises(ValueError):
        mf.make_fixture("x", tmp_path / "raw" / "a.zip", fixtures=tmp_path / "fixtures")
