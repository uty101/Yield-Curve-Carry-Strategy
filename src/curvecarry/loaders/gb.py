"""UK: Bank of England nominal government spot curve, month-end archive (step 1.2).

One zip, ``glcnominalmonthedata.zip``, holding one xlsx per period (1970 to
2015, 2016 to 2024, 2025 to present); the loader reads every xlsx in the zip
and, in each, the sheet whose name contains ``spot curve`` (``4. spot
curve``). Layout: a title row, ``Maturity``, a ``years:`` header row with
the tenors (0.5-year steps, to 25 years in the early file and 40 later),
then one row per month dated at the last business day. Values are percent.

Compounding: the BoE states that its yield curves are *continuously
compounded and quoted on an annual basis* (yield-curves page, FAQ "At which
frequency are the yields compounded", and the EIOPA comparison table where
the annually-compounded line refers to the PRA Solvency II figures, not
these curves). The package convention is annual compounding, so
``z = exp(r) - 1`` is applied here, once (``decisions/compounding.md``).
``curve_type = zero``.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from curvecarry import checks, manifest
from curvecarry.loaders import base

SOURCE = "boe"
COUNTRY = "GB"
CURVE_TYPE = "zero"
NAME = "glcnominalmonthedata"
URL = f"https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/{NAME}.zip"
SHEET_KEY = "spot curve"


def fetch(cfg: dict) -> list[Path]:
    content = manifest.fetch_bytes(URL)
    path = manifest.raw_path(SOURCE, NAME, "zip", base.today())
    manifest.write_raw(content, path, URL, SOURCE)
    return [path]


def continuous_to_annual(r: pd.Series | np.ndarray | float):
    """Continuously compounded decimal rate -> annually compounded decimal rate."""
    return np.exp(r) - 1.0


def _spot_sheet(wb: openpyxl.Workbook):
    names = [n for n in wb.sheetnames if SHEET_KEY in n.lower()]
    if len(names) != 1:
        raise ValueError(f"expected one sheet containing {SHEET_KEY!r}, found {names}")
    return wb[names[0]]


def parse_workbook(data: bytes | Path) -> pd.DataFrame:
    """One xlsx -> long frame ``obs_date, tenor_years, yield`` (decimal, annual compounding)."""
    src = io.BytesIO(data) if isinstance(data, bytes) else data
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = _spot_sheet(wb)
    tenors: list[float] | None = None
    records = []
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        if isinstance(row[0], str) and row[0].strip().lower().startswith("years"):
            tenors = [float(x) if x is not None else np.nan for x in row[1:]]
            continue
        if isinstance(row[0], datetime) and tenors is not None:
            vals = row[1 : 1 + len(tenors)]
            for t, v in zip(tenors, vals, strict=False):
                if v is None or np.isnan(t):
                    continue
                records.append((row[0], t, float(v)))
    if tenors is None or not records:
        raise ValueError("no 'years:' header or no dated rows in the spot curve sheet")
    df = pd.DataFrame(records, columns=["obs_date", "tenor_years", "pct"])
    df["yield"] = continuous_to_annual(
        df["pct"] / 100.0
    )  # percent -> decimal, continuous -> annual
    return df[["obs_date", "tenor_years", "yield"]]


def parse(paths: list[Path]) -> pd.DataFrame:
    """Zip(s) of xlsx and/or bare xlsx -> long monthly frame."""
    frames = []
    for p in paths:
        p = Path(p)
        if p.suffix.lower() == ".zip":
            with zipfile.ZipFile(p) as zf:
                members = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
                if not members:
                    raise ValueError(f"{p}: no xlsx inside")
                for n in members:
                    frames.append(parse_workbook(zf.read(n)))
        elif p.suffix.lower() == ".xlsx":
            frames.append(parse_workbook(p))
        else:
            raise ValueError(f"{p}: expected .zip or .xlsx")
    return pd.concat(frames, ignore_index=True)


def to_monthly(daily_like: pd.DataFrame) -> pd.DataFrame:
    """The archive is already monthly: stamp each row to its calendar month end."""
    df = daily_like.copy()
    df["date"] = base.month_end(df["obs_date"])
    dup = df.duplicated(["date", "tenor_years"], keep=False)
    if dup.any():
        raise ValueError(f"{dup.sum()} rows share a month and tenor; the archive is not monthly")
    return df[["date", "tenor_years", "yield"]]


def latest_paths() -> list[Path]:
    return [manifest.latest_raw(SOURCE, NAME)]


def load(cfg: dict) -> pd.DataFrame:
    monthly = to_monthly(parse(latest_paths()))
    df = base.validate_curve(base.to_curveframe(monthly, COUNTRY, CURVE_TYPE, SOURCE))
    base.write_interim(df, COUNTRY)
    checks.write_coverage(df, "observed")
    return df
