"""Trim a raw download to a committed test fixture.

    uv run python scripts/make_fixture.py <source> <raw_path> [--sheet NAME] [--header-lines N]

Writes the header rows plus the first 50 data rows of ``raw_path`` to
``tests/fixtures/<source>/<name>.<ext>``, where ``<name>`` is the raw
filename with its ``_YYYYMMDD`` fetch date removed (``DGS10_20260916.csv``
→ ``DGS10.csv``). CSV and xlsx are supported; for xlsx ``--sheet`` names the
sheet to keep (the first sheet if omitted) and the output workbook has that
one sheet under the same name. ``--header-lines`` is the number of leading
rows that are header, not data (default 1; MOF has 2, the Bundesbank API
several). Tests read fixtures only (global rule 12); fetches never run in
a test.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

N_DATA_ROWS = 50
FIXTURES = Path("tests/fixtures")
_DATED = re.compile(r"^(?P<name>.+)_\d{8}$")


def fixture_name(raw: Path) -> str:
    m = _DATED.match(raw.stem)
    return (m["name"] if m else raw.stem) + raw.suffix


def trim_csv(raw: Path, out: Path, header_lines: int) -> int:
    with open(raw, encoding="utf-8-sig", newline="") as fh:
        lines = fh.readlines()
    keep = lines[: header_lines + N_DATA_ROWS]
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="") as fh:
        fh.writelines(keep)
    return max(len(keep) - header_lines, 0)


def trim_xlsx(raw: Path, out: Path, header_lines: int, sheet: str | None) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    new = openpyxl.Workbook()
    dst = new.active
    dst.title = ws.title
    n = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= header_lines + N_DATA_ROWS:
            break
        dst.append(list(row))
        n += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    new.save(out)
    return max(n - header_lines, 0)


def make_fixture(
    source: str,
    raw: Path,
    *,
    sheet: str | None = None,
    header_lines: int = 1,
    fixtures: Path = FIXTURES,
) -> Path:
    raw = Path(raw)
    out = Path(fixtures) / source / fixture_name(raw)
    if raw.suffix.lower() == ".csv":
        trim_csv(raw, out, header_lines)
    elif raw.suffix.lower() == ".xlsx":
        trim_xlsx(raw, out, header_lines, sheet)
    else:
        raise ValueError(f"unsupported fixture type {raw.suffix!r}: csv or xlsx only")
    return out


def main(argv: list[str] | None = None) -> None:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("source")
    p.add_argument("raw_path", type=Path)
    p.add_argument("--sheet", default=None, help="xlsx sheet name to keep (default: first)")
    p.add_argument("--header-lines", type=int, default=1, help="leading header rows (default 1)")
    a = p.parse_args(argv)
    out = make_fixture(a.source, a.raw_path, sheet=a.sheet, header_lines=a.header_lines)
    print(out)


if __name__ == "__main__":
    main()
